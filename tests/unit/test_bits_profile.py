"""
ADR-0025 acceptance tests: the document grammar is a BITS profile with an
orthogonal `binding` axis.

Two INSTANCES must validate through the real instantiator and stay green:
the shipped template (byte-for-byte compatible through the presets) and the
reference report's faithful structure (docs/reference/…faithful.yaml, derived
from the NIEHS-10 docx).  The difference between them is a difference in
bindings, never in what the grammar can express.  The remaining tests pin the
resolution rules: type-or-(role, binding), role assertion, binding choice, the
file-only bindings, and the tree carrying both axes.
"""

from pathlib import Path

import pytest
import yaml

from document_model.document_config import _tree_from_document_list, load_template, ACTIVE_TEMPLATE
from document_model.document_template import instantiate
from document_model.document_tree import walk_tree
from document_model.render_capabilities import (
    BINDINGS, COMPONENT_CATALOG, ROLE_PROFILE, preset_for, role_for,
)

REPO = Path(__file__).resolve().parents[2]
REFERENCE_YAML = REPO / "docs" / "reference" / "niehs-10-structure.faithful.yaml"


def _nodes(tree):
    out = []
    walk_tree(tree, out.append)
    return out


# ---------------------------------------------------------------------------
# The two acceptance instances
# ---------------------------------------------------------------------------

def test_shipped_template_is_an_instance_of_the_profile():
    tree = _tree_from_document_list(load_template(ACTIVE_TEMPLATE))
    for node in _nodes(tree):
        assert node.role in ROLE_PROFILE, node.id
        assert node.binding in BINDINGS, node.id


def test_reference_report_is_an_instance_of_the_profile():
    """The NIEHS-10 structure (nested front matter, structured appendices with
    sub-sections two deep, data + authored tables, figures, per-appendix lists,
    55 supplementary files) validates unchanged."""
    document = yaml.safe_load(REFERENCE_YAML.read_text(encoding="utf-8"))["document"]
    tree = _tree_from_document_list(document)
    nodes = _nodes(tree)
    types = {n.node_type for n in nodes}
    # The constructs that used to be violations are all present.
    assert {"appendix", "figure", "data-table", "authored-table", "figures-list",
            "supplementary-material", "freeform-block"} <= types
    about = next(n for n in nodes if n.id == "about-this-report")
    assert [c.title for c in about.children] == ["Authors", "Contributors"]
    c22 = next(n for n in nodes if n.id == "c-2-2")
    assert c22.level == 3 and {c.node_type for c in c22.children} == {"figure", "data-table"}
    assert sum(1 for n in nodes if n.node_type == "supplementary-material") == 55
    # Honest bindings.
    by_id = {n.id: n for n in nodes}
    assert by_id["abstract"].binding == "llm"
    assert by_id["references"].binding == "derived"
    assert by_id["table-d-1"].binding == "authored"
    assert by_id["table-b-1"].binding == "programmatic"
    assert by_id["appendix-a"].binding == "container" and by_id["appendix-a"].role == "app"


# ---------------------------------------------------------------------------
# Resolution rules
# ---------------------------------------------------------------------------

def test_tree_carries_role_and_default_binding():
    tree = instantiate([{"id": "s", "type": "heading-only", "title": "S", "children": [
        {"id": "t", "type": "table", "title": "T", "platform": "Body Weight"}]}])
    assert (tree[0].role, tree[0].binding) == ("sec", "container")
    assert (tree[0].children[0].role, tree[0].children[0].binding) == ("table-wrap", "programmatic")


def test_explicit_role_and_binding_resolve_to_a_preset_without_type():
    tree = instantiate([{"id": "a", "role": "app", "binding": "container", "title": "A"}])
    assert tree[0].node_type == "appendix"
    assert preset_for("table-wrap", "authored") == "authored-table"


def test_explicit_pair_without_a_preset_is_rejected():
    with pytest.raises(ValueError, match="no catalog preset renders role 'toc' with binding 'llm'"):
        instantiate([{"id": "x", "role": "toc", "binding": "llm", "title": "X"}])


def test_role_must_match_the_preset():
    with pytest.raises(ValueError, match="role 'fig' does not match type 'table'"):
        instantiate([{"id": "t", "type": "table", "role": "fig", "title": "T", "platform": "P"}])


def test_binding_must_be_one_the_preset_admits():
    with pytest.raises(ValueError, match="binding 'llm' is not one type 'table' admits"):
        instantiate([{"id": "t", "type": "table", "binding": "llm", "title": "T", "platform": "P"}])


def test_binding_choice_is_kept_on_the_node():
    tree = instantiate([{"id": "r", "type": "narrative", "title": "R", "data_key": "references",
                         "binding": "derived"}])
    assert tree[0].binding == "derived"


def test_containment_is_by_role_not_type():
    # A figure inside a freeform (authored) section inside an appendix: legal
    # because sec → fig and app → sec, with no per-type rule anywhere.
    tree = instantiate([{"id": "a", "type": "appendix", "title": "A", "children": [
        {"id": "s", "type": "freeform-block", "title": "S", "content": {"html": "<p>x</p>"},
         "children": [{"id": "f", "type": "figure", "title": "F", "subtype": "chart", "data_key": "k"}]}]}])
    assert tree[0].children[0].children[0].role == "fig"
    assert role_for("appendix") == "app"


def test_supplementary_material_requires_a_file_and_nothing_else():
    ok = instantiate([{"id": "s", "type": "supplementary-material", "title": "S", "content_file": "a.xlsx"}])
    assert ok[0].binding == "authored" and ok[0].level == 0
    with pytest.raises(ValueError, match="requires `content_file`"):
        instantiate([{"id": "s", "type": "supplementary-material", "title": "S"}])
    with pytest.raises(ValueError, match="only `content_file` applies"):
        instantiate([{"id": "s", "type": "supplementary-material", "title": "S",
                      "content_file": "a.xlsx", "representation": "html"}])


def test_authored_figure_subtypes_name_their_image_file():
    ok = instantiate([{"id": "f", "type": "figure", "title": "F", "subtype": "diagram",
                       "content_file": "reference/flow.png"}])
    assert ok[0].content_file == "reference/flow.png"
    with pytest.raises(ValueError, match="requires `content_file`"):
        instantiate([{"id": "f", "type": "figure", "title": "F", "subtype": "diagram"}])
    with pytest.raises(ValueError, match="figure subtype 'sketch' is not one of"):
        instantiate([{"id": "f", "type": "figure", "title": "F", "subtype": "sketch"}])


def test_every_preset_pair_round_trips_through_preset_for():
    """The canonical preset for (role, default binding) is a type with that
    role admitting that binding — the reverse map is total over the catalog."""
    for name, comp in COMPONENT_CATALOG.items():
        canonical = preset_for(comp.role, comp.bindings[0])
        assert canonical is not None
        c = COMPONENT_CATALOG[canonical]
        assert c.role == comp.role and comp.bindings[0] in c.bindings


# ---------------------------------------------------------------------------
# Appendix-scoped numbering (ADR-0025 §5)
# ---------------------------------------------------------------------------

def _scoped_tree():
    """Body: two tables + a figure.  Appendix A: a table + two figures.
    Appendix B: a table nested two sections deep."""
    return instantiate([
        {"region": "body", "children": [
            {"id": "s", "type": "heading-only", "title": "S", "children": [
                {"id": "t1", "type": "table", "title": "T1", "platform": "Body Weight"},
                {"id": "t2", "type": "data-table", "title": "T2", "data_key": "m"},
                {"id": "f1", "type": "figure", "title": "F1", "subtype": "chart", "data_key": "c"}]}]},
        {"region": "back", "children": [
            {"id": "app-a", "type": "appendix", "title": "A", "children": [
                {"id": "ta", "type": "authored-table", "title": "TA", "content": {"html": "<table/>"}},
                {"id": "fa1", "type": "figure", "title": "FA1", "subtype": "chart", "data_key": "c"},
                {"id": "fa2", "type": "figure", "title": "FA2", "subtype": "diagram", "content_file": "x.png"}]},
            {"id": "app-b", "type": "appendix", "title": "B", "children": [
                {"id": "sb", "type": "freeform-block", "title": "SB", "content": {"html": "<p/>"}, "children": [
                    {"id": "sbb", "type": "freeform-block", "title": "SBB", "content": {"html": "<p/>"}, "children": [
                        {"id": "tb", "type": "data-table", "title": "TB", "data_key": "m"}]}]}]}]},
    ])


def test_appendix_tables_and_figures_number_per_appendix():
    from document_model.document_tree import compute_table_numbers, find_node
    tree = _scoped_tree()
    compute_table_numbers(tree)
    labels = {i: (find_node(i, tree).table_number, find_node(i, tree).table_label)
              for i in ("t1", "t2", "ta", "tb")}
    assert labels == {"t1": (1, "1"), "t2": (2, "2"), "ta": (1, "A-1"), "tb": (1, "B-1")}
    figs = {i: (find_node(i, tree).figure_number, find_node(i, tree).figure_label)
            for i in ("f1", "fa1", "fa2")}
    assert figs == {"f1": (1, "1"), "fa1": (1, "A-1"), "fa2": (2, "A-2")}
    # Scope is recorded on every node inside an appendix, and nowhere else.
    assert find_node("sbb", tree).appendix_scope == "B"
    assert find_node("app-a", tree).appendix_scope is None
    assert find_node("s", tree).appendix_scope is None


def test_scoped_captions_and_prefixes_use_the_label():
    from document_model.document_tree import compute_table_numbers, find_node
    from rendering.render_common import table_caption, figure_prefix
    tree = _scoped_tree()
    compute_table_numbers(tree)
    assert table_caption(find_node("tb", tree), "Rows") == "Table B-1. Rows"
    assert table_caption(find_node("t2", tree), "Rows") == "Table 2. Rows"
    assert figure_prefix(find_node("fa2", tree)) == "Figure A-2. "
    assert figure_prefix(find_node("f1", tree)) == "Figure 1. "


def test_genomics_numbers_continue_the_body_sequence_only():
    """Data-driven genomics tables/charts continue from the last BODY number;
    appendix-scoped numbers must not push them."""
    from document_model.document_tree import (
        compute_table_numbers, assign_genomics_table_numbers, assign_genomics_figure_numbers,
    )
    tree = instantiate([
        {"region": "body", "children": [
            {"id": "t1", "type": "data-table", "title": "T1", "data_key": "m"},
            {"id": "f1", "type": "figure", "title": "F1", "subtype": "chart", "data_key": "c"},
            {"id": "g", "type": "genomics-section", "title": "G", "data_key": "genomics_sections",
             "narrative_key": "gene_set_narrative"}]},
        {"region": "back", "children": [
            {"id": "app-a", "type": "appendix", "title": "A", "children": [
                {"id": "ta1", "type": "data-table", "title": "A1", "data_key": "m"},
                {"id": "ta2", "type": "data-table", "title": "A2", "data_key": "m"},
                {"id": "ta3", "type": "data-table", "title": "A3", "data_key": "m"},
                {"id": "fa1", "type": "figure", "title": "FA1", "subtype": "chart", "data_key": "c"},
                {"id": "fa2", "type": "figure", "title": "FA2", "subtype": "chart", "data_key": "c"}]}]},
    ])
    compute_table_numbers(tree)
    sections = [{"type": "gene_set", "organ": "Liver", "sex": "Male",
                 "charts": [{"type": "umap"}]}]
    assign_genomics_table_numbers(tree, sections)
    assign_genomics_figure_numbers(tree, sections)
    assert sections[0]["table_number"] == 2      # after body Table 1, not after A-3
    assert sections[0]["charts"][0]["figure_number"] == 2


def test_reference_report_labels_match_the_docx():
    """The reference instance numbers exactly as the printed report does."""
    from document_model.document_tree import compute_table_numbers, find_node
    document = yaml.safe_load(REFERENCE_YAML.read_text(encoding="utf-8"))["document"]
    tree = _tree_from_document_list(document)
    compute_table_numbers(tree)
    got = {i: find_node(i, tree).table_label
           for i in ("table-1-sample-counts", "table-b-1", "table-c-1", "table-d-1")}
    assert got == {"table-1-sample-counts": "1", "table-b-1": "B-1", "table-c-1": "C-1", "table-d-1": "D-1"}
    assert [find_node(f"figure-c-{n}", tree).figure_label for n in range(1, 7)] == [f"C-{n}" for n in range(1, 7)]
    assert [find_node(f"figure-d-{n}", tree).figure_label for n in (1, 2)] == ["D-1", "D-2"]
    assert find_node("apical-endpoint-benchmark-dose-summary", tree).table_label == "8"


# ---------------------------------------------------------------------------
# Scoped lists (ADR-0025 phase 4): Contents / Tables / Figures per appendix
# ---------------------------------------------------------------------------

def _listed_tree():
    return instantiate([
        {"region": "front", "children": [
            {"id": "toc", "type": "toc", "title": "Contents"},
            {"id": "tables", "type": "tables-list", "title": "Tables"},
            {"id": "figures", "type": "figures-list", "title": "Figures"}]},
        {"region": "body", "children": [
            {"id": "s", "type": "heading-only", "title": "Results", "children": [
                {"id": "t1", "type": "data-table", "title": "Body table", "data_key": "m"},
                {"id": "f1", "type": "figure", "title": "Body figure", "subtype": "chart", "data_key": "c"}]}]},
        {"region": "back", "children": [
            {"id": "app-a", "type": "appendix", "title": "Alpha", "children": [
                {"id": "toc-a", "type": "toc", "title": "Contents"},
                {"id": "tables-a", "type": "tables-list", "title": "Tables"},
                {"id": "figures-a", "type": "figures-list", "title": "Figures"},
                {"id": "sa", "type": "freeform-block", "title": "Alpha section", "content": {"html": "<p/>"}, "children": [
                    {"id": "ta", "type": "data-table", "title": "Alpha table", "data_key": "m"},
                    {"id": "fa", "type": "figure", "title": "Alpha figure", "subtype": "chart", "data_key": "c"}]}]}]},
    ])


def test_entries_carry_scope_and_figures_are_walked():
    from rendering.report_data_toc import _build_toc_entries
    tree = _listed_tree()
    toc, tables, figures = _build_toc_entries({"m": {"headers": ["h"], "rows": [["1"]]}}, tree=tree)
    assert [(e["title"], e["scope"]) for e in toc] == [
        ("Results", None), ("Appendix A. Alpha", None), ("Alpha section", "A")]
    assert [(e["label"], e["scope"], e["ready"]) for e in tables] == [("1", None, True), ("A-1", "A", True)]
    assert [(e["label"], e["scope"], e["ready"]) for e in figures] == [("1", None, False), ("A-1", "A", False)]


def test_list_nodes_show_only_their_scope_on_every_surface():
    from docx import Document
    from rendering import html_generator as hg, latex_generator as lg, docx_generator as dg
    from rendering.report_data_toc import _build_toc_entries
    from document_model.document_tree import find_node
    tree = _listed_tree()
    toc, tables, figures = _build_toc_entries({}, tree=tree)
    data = {"toc_entries": toc, "table_entries": tables, "figure_entries": figures}
    # Front matter: body only (the appendix title is a body-level placeholder).
    front_tables = hg._render_tables_list(find_node("tables", tree), data)
    assert "Table 1. Body table" in front_tables and "A-1" not in front_tables
    front_toc = hg._render_toc(find_node("toc", tree), data)
    assert "Appendix A. Alpha" in front_toc and "Alpha section" not in front_toc
    assert "Figure 1. Body figure" in hg._render_figures_list(find_node("figures", tree), data)
    # Appendix A's own lists: A-scoped only.
    a_tables = hg._render_tables_list(find_node("tables-a", tree), data)
    assert "Table A-1. Alpha table" in a_tables and "Body table" not in a_tables
    a_toc = hg._render_toc(find_node("toc-a", tree), data)
    assert "Alpha section" in a_toc and "Results" not in a_toc
    # LaTeX: native commands for the body lists, an itemized list per appendix.
    assert lg._render_tables_list(find_node("tables", tree), data).endswith("\\listoftables")
    assert lg._render_toc(find_node("toc", tree), data) == "\\tableofcontents"
    a_tex = lg._render_tables_list(find_node("tables-a", tree), data)
    assert "\\item Table A-1. Alpha table" in a_tex and "Body table" not in a_tex
    assert "\\item Alpha section" in lg._render_toc(find_node("toc-a", tree), data)
    assert "\\item Figure A-1. Alpha figure" in lg._render_figures_list(find_node("figures-a", tree), data)
    # Word: plain entries for the scoped lists.
    doc = Document(); dg._render_tables_list(doc, find_node("tables-a", tree), data)
    assert any("Table A-1. Alpha table" in p.text for p in doc.paragraphs)
    doc = Document(); dg._render_toc(doc, find_node("toc-a", tree), data)
    assert any(p.text == "Alpha section" for p in doc.paragraphs)


def test_appendix_tables_stay_out_of_the_front_latex_list():
    from rendering import latex_generator as lg
    from document_model.document_tree import find_node
    tree = _listed_tree()
    from document_model.document_tree import compute_table_numbers
    compute_table_numbers(tree)
    assert "\\niehsunlistedtrue" in lg._render_appendix(find_node("app-a", tree), {})
    cls = (REPO / "latex" / "niehs.cls").read_text(encoding="utf-8")
    assert "\\ifniehsunlisted\\caption[]{#2}\\else\\caption{#2}\\fi" in cls


def test_reference_lists_per_appendix_match_the_docx():
    from rendering.report_data_toc import _build_toc_entries
    document = yaml.safe_load(REFERENCE_YAML.read_text(encoding="utf-8"))["document"]
    tree = _tree_from_document_list(document)
    toc, tables, figures = _build_toc_entries({}, tree=tree)
    assert [e["label"] for e in tables if e["scope"] is None] == [str(n) for n in range(1, 9)]
    assert [e["label"] for e in figures if e["scope"] == "C"] == [f"C-{n}" for n in range(1, 7)]
    assert [e["label"] for e in tables if e["scope"] == "D"] == ["D-1"]
    assert [e["title"] for e in figures if e["scope"] == "D"][0].startswith("Benchmark Dose Model")
    assert [e["title"] for e in toc if e["scope"] == "C"] == [
        "Gene Expression Quality Control", "Empirical False Discovery Rate", "Methods", "Results"]
