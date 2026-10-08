"""tooling.cebs_to_pool — CEBS Individual Animal workbooks → pool truth CSVs.

The supplied PFHxSAm package's truth CSVs were hand-built from the NIEHS-10
workbooks; the converter must regenerate them byte-for-byte from the published
CEBS workbooks (tests/fixtures/reference/NIEHS-10), so a CEBS package ingests
the same way the supplied one did.  The 2025 (DTT) schema has no Sex column; the sex is
filled from the Key sheet's study title, and an ambiguous title is refused.
"""

from pathlib import Path

import openpyxl
import pytest

from tooling.cebs_to_pool import convert_workbook, key_sheet_sex, workbook_stem

from tests.conftest import GIVEN_PFHXSAM_DIR, REFERENCE_DIR


def test_reference_workbooks_regenerate_truth_csvs_byte_for_byte(tmp_path):
    written = []
    for xlsx in sorted((REFERENCE_DIR / "NIEHS-10").glob("*Individual_Animal_*.xlsx")):
        summary = convert_workbook(str(xlsx), str(tmp_path))
        assert summary["status"] == "ok", summary
        assert summary["sex_filled"] is None  # NIEHS-10 workbooks carry a Sex column
        written += summary["written"]
    assert len(written) == 14
    for name in written:
        ref = GIVEN_PFHXSAM_DIR / "_originals" / name
        if not ref.exists():  # clinical obs CSVs stay in files/ (never pivoted)
            ref = GIVEN_PFHXSAM_DIR / "files" / name
        assert (tmp_path / name).read_bytes() == ref.read_bytes(), name


def _dtt_workbook(path: Path, title: str) -> None:
    """A minimal 2025-schema workbook: no Sex column, 'Study Number'."""
    wb = openpyxl.Workbook()
    key = wb.active
    key.title = "Key to Column Labels"
    key.append(["Column Label", "Explanation of Column Label", "", title])
    data = wb.create_sheet("Data")
    data.append(["Study Number", "Concentration", "Animal ID", "Selection",
                 "Body Weight", "Observation Day", "Terminal Flag"])
    data.append(["108020004", "0", "1", "Core Animals", "180.0", "SD0", ""])
    data.append(["108020004", "0", "1", "Core Animals", "185.5", "SD5", "Yes"])
    wb.save(path)


def test_dtt_workbook_fills_sex_from_key_title(tmp_path):
    src = tmp_path / "108020004_Individual_Animal_Body_Weight_Data_Rats.xlsx"
    _dtt_workbook(src, "5-Day Whole Body Inhalation Toxicity Study of 1,4-Dichlorobenzene "
                       "in Female Sprague Dawley Rats (HSD:Sprague Dawley SD)")
    out = tmp_path / "out"
    out.mkdir()
    summary = convert_workbook(str(src), str(out))
    assert summary["status"] == "ok"
    assert summary["sex_filled"] == "Female"
    lines = (out / "body_weight_truth_female.csv").read_text().splitlines()
    assert lines[0] == ("NTP Study Number,Concentration,Animal ID,Sex,Selection,"
                        "Body Weight,Observation Day,Terminal Flag")
    assert lines[1] == "108020004,0,1,Female,Core Animals,180.0,SD0,"


def test_dtt_workbook_with_ambiguous_title_is_refused(tmp_path):
    src = tmp_path / "X_Individual_Animal_Body_Weight_Data.xlsx"
    _dtt_workbook(src, "Study in Male and Female Rats")
    summary = convert_workbook(str(src), str(tmp_path))
    assert summary["status"] == "error"


def test_key_sheet_sex_reads_only_the_title(tmp_path):
    src = tmp_path / "x.xlsx"
    _dtt_workbook(src, "Study in Female B6D2F1/Crl Mice")
    assert key_sheet_sex(openpyxl.load_workbook(src)) == "Female"


@pytest.mark.parametrize("name,stem", [
    ("108020004_Individual_Animal_Organ_Weight_Data_Rats.xlsx", "organ_weights"),
    ("C20022-01_Individual_Animal_Clinical_Chemistry_Data.xlsx", "clin_chem"),
    ("C20022-01_Individual_Animal_Clinical_Observations_Data.xlsx", "clinical_obs"),
    ("108020004_Individual_Animal_Removal_Reasons_Data_Rats.xlsx", None),
])
def test_workbook_stem(name, stem):
    assert workbook_stem(name) == stem


def test_genomic_metadata_matches_samples_and_uses_vocabulary(tmp_path):
    """Organ comes from the expression file holding the experiment's samples;
    values are spelled in the metadata vocabulary (lowercase organ/sex, common
    species name, S1500+_<species> platform)."""
    import json
    from tooling.cebs_to_pool import genomic_metadata
    (tmp_path / "X_BMDExpress_Expression_Data_Liver_Rats.txt").write_text(
        "Probe\tS1\tS2\nA_1\t1\t2\n")
    (tmp_path / "X_BMDExpress_Expression_Data_Kidney_Rats.txt").write_text(
        "Probe\tS3\tS4\nA_1\t1\t2\n")
    chip = {"provider": "BioSpyder", "species": "Rattus norvegicus", "name": "S1500_Plus_Rat"}
    project = {"doseResponseExperiments": [
        {"name": "rat_4D_LI_No0", "chip": chip, "treatments": [{"name": "S1"}, {"name": "S2"}]},
        {"name": "rat_4D_KI_No0", "chip": chip, "treatments": [{"name": "S3"}, {"name": "S4"}]},
    ]}
    pj = tmp_path / "p.json"
    pj.write_text(json.dumps(project))
    meta = genomic_metadata(str(pj), str(tmp_path), "Female")
    assert meta["rat_4D_LI_No0"] == {
        "organ": "liver", "sex": "female", "species": "rat", "provider": "BioSpyder",
        "platform": "S1500+_rat", "subjectType": "in vivo", "dataType": "gene_expression",
    }
    assert meta["rat_4D_KI_No0"]["organ"] == "kidney"


def test_molar_doses_map_to_the_projects_mg_per_kg_levels():
    """NIEHS-02 records apical doses in mmol/kg; its genomic project and the
    published tables use mg/kg (EHDP 362.4 g/mol: 0.169 → 61.2)."""
    import pytest

    from tooling.cebs_to_pool import molar_to_mg_per_kg

    mapping, ratio = molar_to_mg_per_kg(
        ["0", "0.169", "0.338", "0.675", "1.35", "2.70", "0.169"],
        [61.2, 122.5, 244.6, 489.2, 978.5])
    assert mapping == {"0": "0", "0.169": "61.2", "0.338": "122.5", "0.675": "244.6",
                       "1.35": "489.2", "2.70": "978.5"}
    assert 360 < ratio < 365
    with pytest.raises(ValueError):            # no single molecular weight
        molar_to_mg_per_kg(["0.169", "0.338"], [61.2, 300.0])
    with pytest.raises(ValueError):            # level counts differ
        molar_to_mg_per_kg(["0.169", "0.338"], [61.2])
