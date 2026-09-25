"""
ADR-0025 phase 3: the two authored-content channels the reference report needs.

  - An AUTHORED FIGURE (subtype diagram / photograph) names its image file in
    `content_file`; the shared figure_payload extract reads it from templates/
    and every surface places it (or shows a visible pending note when the file
    is missing).  The LaTeX bundle ships the bytes under figures/.
  - An AUTHORED TABLE whose supplied markup is an HTML <table> becomes a real
    grid on the Word and BITS surfaces (and a tabular on LaTeX when no LaTeX
    source was given) through the shared authored_table_matrix extract.
"""

import base64

from docx import Document
from lxml import etree

from document_model import document_template as dt
from document_model.document_node import DocNode
from rendering import html_generator as hg, latex_generator as lg, docx_generator as dg
from rendering import jats_generator as jg
from rendering.latex_export import _collect_figure_files
from rendering.render_common import authored_table_matrix, figure_payload

# A minimal VALID 4x4 PNG (python-docx reads its dimensions).
_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAQAAAAECAIAAAAmkwkpAAAAEElEQVR4nGM8wYAATAxEcQAo3ADQ"
    "QD5s4QAAAABJRU5ErkJggg=="
)


def _diagram(content_file="reference/flow.png") -> DocNode:
    return DocNode(id="fig-d-1", title="Flowchart", node_type="figure",
                   subtype="diagram", content_file=content_file, figure_label="D-1",
                   figure_number=1)


# ---------------------------------------------------------------------------
# Authored figures
# ---------------------------------------------------------------------------

def test_authored_figure_payload_reads_the_template_file(tmp_path, monkeypatch):
    monkeypatch.setattr(dt, "TEMPLATES_DIR", tmp_path)
    (tmp_path / "reference").mkdir()
    (tmp_path / "reference" / "flow.png").write_bytes(_PNG)
    payload = figure_payload(_diagram(), {})
    assert payload["filename"] == "flow.png" and payload["mimetype"] == "image/png"
    assert base64.b64decode(payload["png_b64"]) == _PNG


def test_authored_figure_renders_on_all_surfaces(tmp_path, monkeypatch):
    monkeypatch.setattr(dt, "TEMPLATES_DIR", tmp_path)
    (tmp_path / "reference").mkdir()
    (tmp_path / "reference" / "flow.png").write_bytes(_PNG)
    fig = _diagram()
    html = hg._render_figure(fig, {})
    assert "data:image/png;base64," in html and "Figure D-1. Flowchart" in html
    tex = lg._render_figure(fig, {})
    assert r"\includegraphics" in tex and "figures/flow.png" in tex and "Figure D-1." in tex
    doc = Document(); dg._render_figure(doc, fig, {})
    assert any("image" in r.reltype for r in doc.part.rels.values())
    # The LaTeX bundle ships the bytes the .tex references.
    files = _collect_figure_files({}, [fig])
    assert files == {"flow.png": _PNG}


def test_missing_authored_image_is_a_visible_gap_everywhere(tmp_path, monkeypatch):
    monkeypatch.setattr(dt, "TEMPLATES_DIR", tmp_path)
    fig = _diagram("reference/absent.png")
    assert figure_payload(fig, {}) is None
    assert "pending" in hg._render_figure(fig, {}).lower()
    assert "pending" in lg._render_figure(fig, {}).lower()
    doc = Document(); dg._render_figure(doc, fig, {})
    assert any("pending" in p.text.lower() for p in doc.paragraphs)
    assert _collect_figure_files({}, [fig]) == {}


def test_latex_refuses_formats_includegraphics_cannot_place(tmp_path, monkeypatch):
    monkeypatch.setattr(dt, "TEMPLATES_DIR", tmp_path)
    (tmp_path / "d.svg").write_text("<svg/>")
    fig = _diagram("d.svg")
    assert "unsupported image format" in lg._render_figure(fig, {})
    assert "data:image/svg+xml;base64," in hg._render_figure(fig, {})


def test_data_figure_payload_is_unchanged():
    """A chart figure still reads its pipeline payload at data[data_key]."""
    fig = DocNode(id="f", title="F", node_type="figure", subtype="chart", data_key="k")
    assert figure_payload(fig, {"k": {"png_b64": "AAAA", "filename": "f.png"}})["filename"] == "f.png"
    assert figure_payload(fig, {}) is None
    # Tree data figures are bundled too (they used to be dropped by the bundle).
    b64 = base64.b64encode(_PNG).decode()
    assert _collect_figure_files({"k": {"png_b64": b64, "filename": "f.png"}}, [fig]) == {"f.png": _PNG}


# ---------------------------------------------------------------------------
# Authored tables
# ---------------------------------------------------------------------------

def _authored_table(html=None, latex=None) -> DocNode:
    node = DocNode(id="table-d-1", title="Rules", node_type="authored-table",
                   table_number=1, table_label="D-1")
    node.resolved_content = {"html": html, "latex": latex}
    return node


def test_authored_table_matrix_parses_headers_and_rows():
    node = _authored_table("<p>intro</p><table><thead><tr><th>A</th><th>B</th></tr></thead>"
                           "<tbody><tr><td>1</td><td>x</td></tr><tr><td>2</td><td>y</td></tr></tbody></table>")
    assert authored_table_matrix(node) == {
        "caption": None, "headers": ["A", "B"], "rows": [["1", "x"], ["2", "y"]], "footnotes": []}
    # A bare <th> first row is a header too; no <table> → None.
    assert authored_table_matrix(_authored_table("<table><tr><th>H</th></tr><tr><td>v</td></tr></table>"))["headers"] == ["H"]
    assert authored_table_matrix(_authored_table("<p>no table</p>")) is None
    assert authored_table_matrix(_authored_table(None, "\\begin{tabular}")) is None


def test_authored_html_table_becomes_a_grid_on_word_bits_and_latex():
    node = _authored_table("<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>x_y</td></tr></table>")
    doc = Document(); dg._render_authored_table(doc, node, {})
    assert len(doc.tables) == 1 and [c.text for c in doc.tables[0].rows[1].cells] == ["1", "x_y"]
    assert any("D-1. Rules" in p.text for p in doc.paragraphs)   # "Table" + nbsp precedes
    sec = etree.Element("sec"); jg._emit_authored_table(sec, node, {})
    xml = etree.tostring(sec).decode()
    assert "<label>Table D-1</label>" in xml and "<td>x_y</td>" in xml and "TODO" not in xml
    tex = lg._render_authored_table(node, {})
    assert "\\begin{tabular}" in tex and "x\\_y" in tex and "Table D-1. Rules" in tex


def test_authored_latex_table_is_verbatim_on_latex_and_traced_on_bits():
    node = _authored_table(latex="\\begin{tabular}{l}only\\end{tabular}")
    assert "only\\end{tabular}" in lg._render_authored_table(node, {})
    sec = etree.Element("sec"); jg._emit_authored_table(sec, node, {})
    xml = etree.tostring(sec).decode()
    assert "<label>Table D-1</label>" in xml and "TODO tracer" in xml
