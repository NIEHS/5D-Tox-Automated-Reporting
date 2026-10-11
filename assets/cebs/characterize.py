"""Characterize every file in a CEBS report folder -> <id>/_characterization.json.

Extracts, per file: kind, size, sha256, plus kind-specific facts (BMDExpress
analysis notes/versions/timestamps, docx/xlsx/pdf metadata dates and headings,
zip members).  Zips are extracted to <id>/_unzipped/<zipname>/ and their
members characterized too.
"""
import sys, os, json, hashlib, zipfile, gzip, re, warnings, datetime
warnings.filterwarnings("ignore")

def sha(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""): h.update(b)
    return h.hexdigest()

def iso(d):
    return d.isoformat() if isinstance(d, (datetime.datetime, datetime.date)) else (str(d) if d else None)

def text_head(p, n=12):
    with open(p, errors="replace") as f:
        return [next(f, "").rstrip("\n")[:300] for _ in range(n)]

def bmdx_text(p):
    """BMDExpress text export / expression file: header notes + table shape."""
    lines = open(p, errors="replace").read().splitlines()
    notes = []
    for l in lines[:80]:
        if "\t" in l: break
        if l.strip(): notes.append(l.strip()[:300])
    tab = [l for l in lines if "\t" in l]
    hdr = tab[0].split("\t")[:40] if tab else []
    return dict(notes=notes, n_lines=len(lines), n_tab_rows=len(tab), header=hdr,
                n_cols=len(tab[0].split("\t")) if tab else 0,
                row2=(tab[1].split("\t")[:12] if len(tab) > 1 else None))

def bmdx_json(p):
    J = json.load(open(p))
    out = dict(project=J.get("name"), keys=list(J.keys()))
    exps = []
    for e in J.get("doseResponseExperiments") or []:
        doses = sorted({t.get("dose") for t in e.get("treatments") or []})
        exps.append(dict(name=e.get("name"), n_animals=len(e.get("treatments") or []),
                         n_probes=len(e.get("probeResponses") or []), doses=doses,
                         chip=(e.get("chip") or {}).get("name") if isinstance(e.get("chip"), dict) else e.get("chip"),
                         log=e.get("logTransformation"),
                         notes=(e.get("analysisInfo") or {}).get("notes")))
    out["experiments"] = exps
    for k in ["oneWayANOVAResults", "williamsTrendResults", "curveFitPrefilterResults",
              "oriogenResults", "bMDResult", "categoryAnalysisResults"]:
        rs = []
        for r in J.get(k) or []:
            lists = {kk: len(v) for kk, v in r.items() if isinstance(v, list)}
            rs.append(dict(name=r.get("name"), lists=lists,
                           notes=(r.get("analysisInfo") or {}).get("notes")))
        out[k] = rs
    return out

def docx_info(p):
    import docx
    d = docx.Document(p); cp = d.core_properties
    paras = [x.text.strip() for x in d.paragraphs if x.text.strip()]
    return dict(created=iso(cp.created), modified=iso(cp.modified), author=cp.author,
                last_modified_by=cp.last_modified_by, n_paras=len(paras), n_tables=len(d.tables),
                n_images=len(d.inline_shapes), first_paras=[x[:300] for x in paras[:6]],
                table_headers=[[c.text.strip()[:60] for c in t.rows[0].cells][:12] for t in d.tables[:4]])

def xlsx_info(p):
    import openpyxl
    if not p.lower().endswith((".xlsx", ".xlsm")):          # mislabeled workbook: openpyxl checks the extension
        import tempfile; tmp = os.path.join(tempfile.mkdtemp(), "x.xlsx"); os.symlink(os.path.abspath(p), tmp); p = tmp
    wb = openpyxl.load_workbook(p, read_only=False, data_only=True)
    pr = wb.properties
    sheets = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(min_row=1, max_row=3, values_only=True))
        sheets.append(dict(name=ws.title, dims=f"{ws.max_row}x{ws.max_column}",
                           rows=[[str(c)[:60] for c in r if c is not None][:20] for r in rows]))
    return dict(created=iso(pr.created), modified=iso(pr.modified), creator=pr.creator, sheets=sheets)

def pdf_info(p):
    import pypdf
    r = pypdf.PdfReader(p); m = r.metadata or {}
    t = (r.pages[0].extract_text() or "").strip().splitlines()
    return dict(pages=len(r.pages), created=str(m.get("/CreationDate")), modified=str(m.get("/ModDate")),
                producer=str(m.get("/Producer")), creator=str(m.get("/Creator")), first_lines=[x[:200] for x in t[:10]])

def csv_info(p):
    h = text_head(p, 4); return dict(head=h)

def sniff(p):
    """Actual type from content (magic bytes), independent of the extension."""
    with open(p, "rb") as f: head = f.read(8)
    if head[:2] == b"\xac\xed": return ".bm2"            # Java serialization (BMDExpress project)
    if head[:4] == b"%PDF": return ".pdf"
    if head[:2] == b"\x1f\x8b": return ".gz"
    if head[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1": return ".doc"   # OLE2 (legacy Office)
    if head[:2] == b"PK":
        try:
            names = zipfile.ZipFile(p).namelist()
        except Exception: return ".zip"
        if "word/document.xml" in names: return ".docx"
        if "xl/workbook.xml" in names: return ".xlsx"
        if "ppt/presentation.xml" in names: return ".pptx"
        return ".zip"
    if head[:1] in (b"{", b"["): return ".json"
    return None

def characterize(p, depth=0, root=None):
    rec = dict(path=os.path.relpath(p, root), size=os.path.getsize(p), sha256=sha(p))
    if os.path.basename(p).startswith("~$"):
        rec["kind"] = "office_lock_file"; return rec
    ext = os.path.splitext(p)[1].lower(); actual = sniff(p)
    if actual and actual != ext and not (ext in (".txt", ".tsv", ".csv") and actual is None):
        rec["declared_ext"] = ext; rec["actual_type"] = actual
    low = p.lower() if not actual or actual == ext else p.lower() + actual
    try:
        if low.endswith(".zip"):
            z = zipfile.ZipFile(p)
            rec["kind"] = "zip"
            rec["members"] = [dict(name=i.filename, size=i.file_size,
                                   mtime="%04d-%02d-%02d" % i.date_time[:3]) for i in z.infolist() if not i.is_dir()]
            dest = os.path.join(root, "_unzipped", os.path.basename(p)[:-4])
            if not os.path.isdir(dest): z.extractall(dest)
        elif low.endswith(".json"):
            rec["kind"] = "bmdexpress_json"; rec["bmdx"] = bmdx_json(p)
        elif low.endswith(".bm2"):
            rec["kind"] = "bm2"
            out = os.path.join(root, "_bm2json", os.path.basename(p)[:-4] + ".json")
            os.makedirs(os.path.dirname(out), exist_ok=True)
            if not os.path.exists(out):
                import subprocess
                subprocess.run([os.path.join(os.path.dirname(os.path.abspath(__file__)), "bm2json.sh"),
                                os.path.abspath(p), os.path.abspath(out)], check=True, capture_output=True)
            rec["bmdx"] = bmdx_json(out)
        elif low.endswith(".docx"):
            rec["kind"] = "docx"; rec["docx"] = docx_info(p)
        elif low.endswith(".doc"):
            rec["kind"] = "doc"
        elif low.endswith((".xlsx", ".xlsm")):
            rec["kind"] = "xlsx"; rec["xlsx"] = xlsx_info(p)
        elif low.endswith(".pdf"):
            rec["kind"] = "pdf"; rec["pdf"] = pdf_info(p)
        elif low.endswith((".txt", ".tsv")):
            rec["kind"] = "text"; rec["text"] = bmdx_text(p)
        elif low.endswith(".csv"):
            rec["kind"] = "csv"; rec["csv"] = csv_info(p)
        elif low.endswith(".gz"):
            rec["kind"] = "gz"; rec["head"] = [l[:200] for l in gzip.open(p, "rt", errors="replace").read(4000).splitlines()[:4]]
        else:
            rec["kind"] = os.path.splitext(p)[1].lower() or "other"
    except Exception as e:
        rec["error"] = f"{type(e).__name__}: {e}"[:300]
    return rec

if __name__ == "__main__":
    root = sys.argv[1]
    recs = []
    files = sorted(f for f in os.listdir(root) if os.path.isfile(os.path.join(root, f)) and not f.startswith("_") and not f.endswith(".part"))
    for f in files:
        recs.append(characterize(os.path.join(root, f), root=root))
    # members of extracted zips (skip binaries/code trees)
    uz = os.path.join(root, "_unzipped")
    os.makedirs(uz, exist_ok=True)
    for dp, dn, fn in os.walk(uz):
        if any(s in dp for s in ("/bin", "__pycache__", "/bmds/")): continue
        for f in sorted(fn):
            if f.lower().endswith((".exe", ".py", ".pyc", ".dll", ".jar", ".cel", ".cel.gz", ".png", ".jpg", ".html", ".js", ".css")): continue
            p = os.path.join(dp, f)
            if os.path.getsize(p) > 600e6: continue
            recs.append(characterize(p, depth=1, root=root))
    json.dump(recs, open(os.path.join(root, "_characterization.json"), "w"), indent=1, default=str)
    print(root, len(recs), "records")
