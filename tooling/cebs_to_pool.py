#!/usr/bin/env python3
"""tooling/cebs_to_pool.py — convert a published CEBS NIEHS data package into
pool inputs rlm-bmdx ingests.

CEBS publishes each study's apical data as NTP "Individual Animal" workbooks
(`Key to Column Labels` + `Data` sheets).  The pool's accepted apical source is
the long-format `*_truth_<sex>.csv` (one row per animal/observation, NTP columns):
the filename sets platform + `tox_study`, and on upload `fingerprint_and_store`
converts it into the wide BMDExpress pivot + per-animal `.sidecar.json`.  The
reference set's truth CSVs are exactly the workbook `Data` sheet split by sex —
`tests/unit/test_cebs_to_pool.py` pins that by regenerating them byte-for-byte
from the reference workbooks.

The `Data` headers carry no units; the Key sheet states them per column, and
they differ by study (plasma ng/ml in NIEHS-07 – -09, ug/ml in NIEHS-10).  Each
CSV gets a companion `<stem>_truth_<sex>.units.json` ({column: unit},
bmdx_pipe.key_sheet_units) that the pool's ingest records in the sidecar.

NIEHS-02 … 05 publish the CEBS database-export layout instead of Key/Data
(one row per animal per measurement); tooling.cebs_db_export converts it to
the same Key/Data columns first.

Dose unit: RR-8 and NIEHS-02 … 06 record apical doses in mmol/kg (the Key
sheet's "Dose in mMol/kg", or the export's DOSE_UNIT), while their genomic
BMDExpress project and published tables use mg/kg (EHDP: 0.169 mmol/kg =
61.2 mg/kg).  `--mg-per-kg-from <project .json>` converts a molar dose to the
project's own mg/kg dose level: the project's nonzero levels are aligned with
a run of the molar levels and must share one ratio (the molecular weight,
within 1%), or the workbook is refused; a level the project lacks (NIEHS-05's
top dose) is computed at that ratio.
Without it the doses are left as recorded and the summary reports the unit.

NIEHS-12 and -13 publish blood and tissue concentrations only in the Bookshelf
bundle, as the analytical lab's results workbook (one sheet per species);
tooling.cebs_lab_concentration converts the sheet of the package's species —
named by its sibling workbooks' "_Rats"/"_Mice" suffix — and `--bundle` adds
the bundle directory to the workbooks converted.

The 2025 (DTT) workbooks changed the schema: the `Data` sheet has NO `Sex`
column (the studies are single-sex) and says `Study Number` instead of
`NTP Study Number`.  The sex is stated only in the Key sheet's study title
("… in Female Sprague Dawley Rats …").  We fill the `Sex` column from that title
— an inference from the file's own text, reported in the run summary — and
refuse a workbook whose title names both sexes or neither.

Gene expression: the BMDExpress project's experiment names are tissue codes
(`rat_4D_HE_No0`).  The organ is taken from DATA, not the code: each project
experiment's sample IDs are matched against the package's per-tissue expression
files (`…_Expression_Data_<Tissue>_<Species>.txt`), whose sample columns are the
same IDs.  `--metadata-out` writes the resulting per-experiment
ExperimentDescription map ({name: {organ, sex, species, …}}) for the
`/api/experiment-metadata` approval step.

Usage:
    python -m tooling.cebs_to_pool --package assets/cebs/NIEHS-13 \\
        --prefix 108020004 --out sessions/<id>/files \\
        [--bundle assets/cebs/_bookshelf/NIEHS-13] \\
        [--genomic-bm2-json assets/cebs/NIEHS-13/_bm2json/108020004_4D_Genomic_Rats.json] \\
        [--metadata-out /tmp/metadata.json]
"""

from __future__ import annotations

import argparse
import csv
import glob
import json
import os
import re
from pathlib import Path

import openpyxl
from bmdx_pipe import key_sheet_units

from tooling import cebs_db_export, cebs_lab_concentration

# Individual-animal workbook → truth-CSV stem (the pool's filename patterns:
# bmdx_pipe.file_integrator._PLATFORM_PATTERNS).  Order matters: "Clinical
# Chemistry" must win over "Clinical Observations" only by its own token.
_WORKBOOK_STEMS: list[tuple[re.Pattern, str]] = [
    (re.compile(r"Body_Weight", re.I), "body_weight"),
    (re.compile(r"Organ_Weight", re.I), "organ_weights"),
    (re.compile(r"Clinical_Chemistry", re.I), "clin_chem"),
    (re.compile(r"Clinical_Observations", re.I), "clinical_obs"),
    (re.compile(r"Hematology", re.I), "hematology"),
    (re.compile(r"Hormone", re.I), "hormones"),
    (re.compile(r"Tissue_Concentration", re.I), "tissue_conc"),
]

_SEX_WORDS = re.compile(r"\b(Male|Female)s?\b", re.I)

# Chip species (BMDExpress ChipInfo, Latin) → the metadata vocabulary's common
# name (bmdx_pipe.experiment_metadata).
_CHIP_SPECIES = {"rattus norvegicus": "rat", "mus musculus": "mouse", "homo sapiens": "human"}


def workbook_stem(filename: str) -> str | None:
    """Truth-CSV stem for an Individual Animal workbook, or None (no apical platform)."""
    for pat, stem in _WORKBOOK_STEMS:
        if pat.search(filename):
            return stem
    return None


def _cell(v) -> str:
    """Workbook cell → CSV text, matching how the reference CSVs were written."""
    if v is None:
        return ""
    return str(v)


def key_sheet_sex(wb) -> str | None:
    """The single sex named in the Key sheet's study title, or None.

    Only the title cells of the first row are read — the rest of the Key sheet
    explains column labels (which mention "Sex" generically)."""
    if "Key to Column Labels" not in wb.sheetnames:
        return None
    first = next(wb["Key to Column Labels"].iter_rows(max_row=1, values_only=True), ())
    title = " ".join(str(c) for c in first if c)
    sexes = {m.group(1).title() for m in _SEX_WORDS.finditer(title)}
    return sexes.pop() if len(sexes) == 1 else None


_SPECIES_SUFFIX = {"rats": "rat", "mice": "mouse"}


_DOSE_IN = re.compile(r"\b(?:dose|concentration)\s+in\s+(\S+)", re.I)


def key_sheet_dose_unit(wb) -> str | None:
    """The dose unit the Key sheet states for the Concentration column
    ("Dose in mMol/kg"), or None."""
    if "Key to Column Labels" not in wb.sheetnames:
        return None
    for row in wb["Key to Column Labels"].iter_rows(values_only=True):
        if row and _cell(row[0]).strip() == "Concentration" and len(row) > 1:
            m = _DOSE_IN.search(_cell(row[1]))
            return m.group(1) if m else None
    return None


def project_dose_levels(project_json: str) -> list[float]:
    """The nonzero dose levels of a BMDExpress project's gene-expression
    experiments (mg/kg), ascending."""
    project = json.load(open(project_json))
    levels = set()
    for e in project.get("doseResponseExperiments") or []:
        for t in e.get("treatments") or []:
            if isinstance(t, dict) and t.get("dose"):
                levels.add(float(t["dose"]))
    return sorted(levels)


def molar_to_mg_per_kg(doses: list[str], mg_levels: list[float]) -> tuple[dict[str, str], float, list[str]]:
    """{molar dose text: mg/kg dose text}, the ratio (g/mol), and the molar
    levels the project lacks.

    The project's nonzero levels are aligned with a contiguous run of the
    molar levels (a genomic project may omit a dose — NIEHS-05 has no
    2.7 mmol/kg arrays) and must share one ratio within 1%.  Aligned levels
    take the project's mg/kg value; a level it lacks is computed at the ratio,
    to the project's decimals.  Raises ValueError when no alignment shares a
    ratio."""
    molar = sorted({float(d) for d in doses if d and float(d) > 0})
    k = len(mg_levels)
    if not k or k > len(molar):
        raise ValueError(f"{len(molar)} molar dose level(s) vs {k} in the project")
    for offset in range(len(molar) - k + 1):
        window = molar[offset:offset + k]
        ratios = [m / d for d, m in zip(window, mg_levels)]
        mean = sum(ratios) / k
        if all(abs(r - mean) <= 0.01 * mean for r in ratios):
            break
    else:
        raise ValueError(f"no run of the molar levels {molar} shares one ratio with the "
                         f"project's {mg_levels}")
    decimals = max(len(f"{m:g}".partition(".")[2]) for m in mg_levels)
    by_value = dict(zip(window, mg_levels))
    missing = [d for d in molar if d not in by_value]
    for d in missing:
        by_value[d] = round(d * mean, decimals)
    out = {}
    for d in doses:
        if d and float(d) > 0:
            out[d] = f"{by_value[float(d)]:g}"
        elif d:
            out[d] = "0"
    return out, mean, [f"{d:g}" for d in missing]


def _convert_doses(header, body, unit, mg_levels, notes) -> str | None:
    """Rewrite a molar Concentration column to mg/kg in place; returns the
    unit the doses are now in."""
    if not unit or "mol" not in unit.lower():
        return unit
    if mg_levels is None:
        notes.append(f"doses recorded in {unit}; not converted (no --mg-per-kg-from)")
        return unit
    i = header.index("Concentration")
    mapping, ratio, missing = molar_to_mg_per_kg([_cell(r[i]).strip() for r in body], mg_levels)
    for r in body:
        r[i] = mapping.get(_cell(r[i]).strip(), r[i])
    notes.append(f"doses converted from {unit} to mg/kg by the project's dose levels "
                 f"(×{ratio:.1f} g/mol)"
                 + (f"; {', '.join(missing)} {unit} not in the project, computed at that ratio"
                    if missing else ""))
    return "mg/kg"


def package_species(filenames: list[str]) -> str | None:
    """The one species the workbook names' "_Rats"/"_Mice" suffix states, or None."""
    found = {_SPECIES_SUFFIX[m.group(1).lower()] for f in filenames
             if (m := re.search(r"_(Rats|Mice)\.xlsx$", f, re.I))}
    return found.pop() if len(found) == 1 else None


def convert_workbook(path: str, out_dir: str, species: str | None = None,
                     mg_levels: list[float] | None = None) -> dict:
    """Write one `<stem>_truth_<sex>.csv` per sex.  Returns a summary dict.
    `species` picks the sheet of a multi-species lab results workbook;
    `mg_levels` (project_dose_levels) converts molar doses to mg/kg."""
    name = os.path.basename(path)
    stem = workbook_stem(name)
    if stem is None:
        return {"file": name, "status": "skipped", "reason": "no apical platform"}

    wb = openpyxl.load_workbook(path, data_only=True)
    if cebs_lab_concentration.is_lab_results(wb.sheetnames):
        if species is None:
            return {"file": name, "status": "error", "reason": "lab results workbook: species unknown"}
        sheets = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
        try:
            sheet = cebs_lab_concentration.species_sheet(sheets, species)
            header, body, lab_units, notes = cebs_lab_concentration.to_key_data(
                sheets[sheet], name.split("_", 1)[0])
        except ValueError as e:
            return {"file": name, "status": "error", "reason": str(e)}
        return _write_csvs(name, out_dir, stem, header, body, header.index("Sex"), lab_units,
                           [f"sheet '{sheet}'"] + notes, None)
    if "Data" not in wb.sheetnames:
        # NIEHS-02 … 05: the CEBS database-export layout (tooling.cebs_db_export).
        sheet = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        export_header = [_cell(c) for c in sheet[0]]
        if not cebs_db_export.is_db_export(export_header):
            return {"file": name, "status": "error", "reason": "no Data sheet and not a CEBS database export"}
        export_rows = [r for r in sheet[1:] if any(c is not None and str(c).strip() for c in r)]
        header, body, export_units, notes = cebs_db_export.to_key_data(stem, export_header, export_rows)
        notes = notes + ["CEBS database export converted to Key/Data columns"]
        unit_col = export_header.index("DOSE_UNIT") if "DOSE_UNIT" in export_header else None
        export_dose_units = {_cell(r[unit_col]).strip() for r in export_rows} - {""} \
            if unit_col is not None else set()
        try:
            dose_unit = _convert_doses(header, body, next(iter(export_dose_units), None)
                                       if len(export_dose_units) == 1 else None, mg_levels, notes)
        except ValueError as e:
            return {"file": name, "status": "error", "reason": f"dose conversion: {e}"}
        return _write_csvs(name, out_dir, stem, header, body, header.index("Sex"), export_units,
                           notes, None, dose_unit)
    rows = [list(r) for r in wb["Data"].iter_rows(values_only=True)]
    header = [_cell(c) for c in rows[0]]
    body = [r for r in rows[1:] if any(c is not None and str(c).strip() for c in r)]

    notes = []
    if header and header[0] == "Study Number":
        header[0] = "NTP Study Number"
        notes.append("renamed 'Study Number' → 'NTP Study Number'")

    if "Sex" in header:
        sex_idx = header.index("Sex")
        filled = None
    else:
        filled = key_sheet_sex(wb)
        if filled is None:
            return {"file": name, "status": "error",
                    "reason": "no Sex column and the Key title does not name exactly one sex"}
        animal_idx = header.index("Animal ID")
        sex_idx = animal_idx + 1
        header.insert(sex_idx, "Sex")
        body = [r[:sex_idx] + [filled] + r[sex_idx:] for r in body]
        notes.append(f"Sex column absent — filled '{filled}' from the Key sheet title (inferred)")

    units = {col: u for col, u in key_sheet_units(path).items() if col in header}
    body = [list(r) for r in body]
    try:
        dose_unit = _convert_doses(header, body, key_sheet_dose_unit(wb), mg_levels, notes)
    except ValueError as e:
        return {"file": name, "status": "error", "reason": f"dose conversion: {e}"}
    return _write_csvs(name, out_dir, stem, header, body, sex_idx, units, notes, filled, dose_unit)


def _write_csvs(name, out_dir, stem, header, body, sex_idx, units, notes, filled,
                dose_unit=None) -> dict:
    """One `<stem>_truth_<sex>.csv` (+ `.units.json`) per sex."""
    by_sex: dict[str, list[list]] = {}
    for r in body:
        by_sex.setdefault(_cell(r[sex_idx]).strip().lower(), []).append(r)

    written = []
    for sex, sex_rows in sorted(by_sex.items()):
        out = os.path.join(out_dir, f"{stem}_truth_{sex}.csv")
        with open(out, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f, lineterminator="\r\n")  # the reference CSVs are CRLF
            w.writerow(header)
            for r in sex_rows:
                w.writerow([_cell(c) for c in r[:len(header)]])
        written.append(os.path.basename(out))
        if units:
            Path(out).with_suffix(".units.json").write_text(json.dumps(units, indent=2) + "\n")
    return {"file": name, "status": "ok", "written": written, "notes": notes,
            "sex_filled": filled, "units": len(units), "dose_unit": dose_unit}


def _by_ref(node, out: dict | None = None) -> dict:
    """Every object of a BMDExpress project JSON by its Jackson '@ref' id.

    The project writes a shared object (e.g. one ChipInfo for several
    experiments) in full once and as its bare '@ref' id everywhere after."""
    out = {} if out is None else out
    if isinstance(node, dict):
        if "@ref" in node:
            out[node["@ref"]] = node
        for v in node.values():
            _by_ref(v, out)
    elif isinstance(node, list):
        for v in node:
            _by_ref(v, out)
    return out


def genomic_metadata(bm2_json_path: str, package_dir: str, sex: str | None) -> dict:
    """{experiment name: ExperimentDescription-shaped dict} for a genomic project.

    Organ comes from matching each experiment's sample IDs to a per-tissue
    expression file; species from the experiment's chip."""
    project = json.load(open(bm2_json_path))
    tissue_samples: dict[str, set[str]] = {}
    for f in glob.glob(os.path.join(package_dir, "*Expression_Data_*.txt")):
        m = re.search(r"Expression_Data_([A-Za-z]+)", os.path.basename(f))
        if not m:
            continue
        with open(f, encoding="utf-8", errors="replace") as fh:
            for line in fh:
                if "\t" in line and not line.startswith("#"):
                    tissue_samples[f] = (m.group(1), set(line.rstrip("\n").split("\t")[1:]))
                    break

    refs = _by_ref(project)
    out = {}
    for exp in project.get("doseResponseExperiments", []):
        samples = {t.get("name") for t in exp.get("treatments") or []}
        hits = [(t, len(samples & s)) for t, s in tissue_samples.values() if samples & s]
        hits.sort(key=lambda x: -x[1])
        if not hits or hits[0][1] != len(samples):
            raise ValueError(f"{exp.get('name')}: no expression file holds all of its samples")
        chip = exp.get("chip")
        chip = refs.get(chip, {}) if not isinstance(chip, dict) else chip
        species = _CHIP_SPECIES.get((chip.get("species") or "").lower())
        # S1500_Plus_Rat → "S1500+_rat" (the vocabulary's platform spelling)
        m = re.fullmatch(r"S1500_Plus_(\w+)", chip.get("name") or "")
        out[exp["name"]] = {
            "organ": hits[0][0].lower(), "sex": sex.lower() if sex else None,
            "species": species, "provider": chip.get("provider"),
            "platform": f"S1500+_{m.group(1).lower()}" if m else chip.get("name"),
            "subjectType": "in vivo", "dataType": "gene_expression",
        }
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--package", required=True, help="CEBS package dir (assets/cebs/<ID>)")
    ap.add_argument("--prefix", default="", help="study-number filename prefix to select one species")
    ap.add_argument("--out", required=True, help="output dir (a session's files/)")
    ap.add_argument("--bundle", help="Bookshelf bundle dir with more Individual Animal workbooks")
    ap.add_argument("--genomic-bm2-json")
    ap.add_argument("--metadata-out")
    ap.add_argument("--mg-per-kg-from", help="the genomic BMDExpress project (.json) whose mg/kg "
                    "dose levels a molar workbook dose is converted to")
    a = ap.parse_args()
    mg_levels = project_dose_levels(a.mg_per_kg_from) if a.mg_per_kg_from else None

    os.makedirs(a.out, exist_ok=True)
    sexes = set()
    pattern = f"{a.prefix}*Individual_Animal_*.xlsx"
    paths = sorted(glob.glob(os.path.join(a.package, pattern)))
    species = package_species(paths)
    if a.bundle:
        paths += sorted(glob.glob(os.path.join(a.bundle, pattern)))
    failed = 0
    for path in paths:
        s = convert_workbook(path, a.out, species, mg_levels)
        print(json.dumps(s))
        failed += s.get("status") == "error"
        if s.get("status") == "ok":
            sexes.update(w.rsplit("_", 1)[1][:-4] for w in s["written"])
    if a.genomic_bm2_json and a.metadata_out:
        study_sex = sexes.pop().title() if len(sexes) == 1 else None
        meta = genomic_metadata(a.genomic_bm2_json, a.package, study_sex)
        Path(a.metadata_out).write_text(json.dumps(meta, indent=2))
        print(json.dumps({"metadata": {k: v["organ"] for k, v in meta.items()}, "sex": study_sex}))
    if failed:
        raise SystemExit(f"{failed} workbook(s) failed to convert")


if __name__ == "__main__":
    main()
