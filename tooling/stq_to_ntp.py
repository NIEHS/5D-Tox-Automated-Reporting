#!/usr/bin/env python3
"""tooling/stq_to_ntp.py — convert raw STQ LIMS body/organ-weight extracts into
the long-format NTP "truth" CSVs the rlm-bmdx pool ingests.

The Triclosan/BPAF apical weights arrived as raw study-database exports
(`assets/IAD_STQ_{BODYWEIGHTS,ORGANWEIGHTS}_*.xlsx`): DB-native columns, a `SQL`
query sheet, both chemicals combined by `DEPOSITOR_STUDY_NUMBER`, and organ
weights in LONG form (one row per animal×organ).

The pool's accepted apical source is a long-format `*_truth*.csv` (one row per
animal/observation, NTP columns): the filename sets platform + `tox_study`
(file_integrator patterns), and on upload `fingerprint_and_store` converts it
into the wide `{stem}.txt` BMDExpress pivot + `{stem}.sidecar.json` per-animal
metadata (bmdx_pipe.tox_study_csv_to_pivot_txt) — exactly how the reference
`body_weight_truth_male.*` / `organ_weights_truth_male.*` were produced.  (xlsx
study files are NOT an accepted source: the BMDProject schema rejects tier xlsx.)

Study→chemical (confirmed by matching each STQ study's dose series to the
session's genomics doses):
  C08002-01  → BPAF      (DTXSID7037717)   doses 0/8/16/31.25/62.5/125/250/500/1000
  C03039C-01 → Triclosan (DTXSID5032498)   doses 0/7.8/15.6/31.25/62.5/125/250/500/1000

Assumptions (all from the delivered data, flagged deliberately):
  * The source is MALE-ONLY (STQ `SEX` is `Male` throughout).
  * `SELECTION` (`None` / `Core Males`) maps to `Core Animals` — no biosampling
    cohort is present in the extract, so every animal is treated as core-study.
  * `DOSE_FACTOR` is empty, so `Concentration = DOSE` (mg/kg) directly.
  * Organ weights are pivoted to one row per animal (Kidney-Left / Kidney-Right /
    Liver / Terminal Body Weight); the reference's `Heart` column is absent here
    and is simply not emitted (the studies did not weigh heart).

Output: one folder per chemical (both produce the same filenames):
    <out>/<Chem>/body_weight_truth_male.csv
    <out>/<Chem>/organ_weights_truth_male.csv

Usage:
    python -m tooling.stq_to_ntp [--body PATH] [--organ PATH] [--out DIR]
"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

import openpyxl

# NTP study-number pattern the pool's fingerprint extractor accepts.
_STUDY_RE = re.compile(r"^[A-Z]\d+-\d+$")


def _normalize_study(raw: str) -> str:
    """Return a study number in the NTP `C12345-01` form.

    The extract mixes two spellings (`C03039C-01` in the organ file, `C03039-01`
    in the body-weight filename); strip the trailing study-variant letter so both
    weight files for a chemical carry one study number."""
    s = str(raw).strip()
    if _STUDY_RE.match(s):
        return s
    m = re.match(r"^([A-Z]\d+)[A-Z]+(-\d+)$", s)
    return f"{m.group(1)}{m.group(2)}" if m else s


STUDY_TO_CHEM = {
    "C08002-01": ("BPAF", "DTXSID7037717"),
    "C03039C-01": ("Triclosan", "DTXSID5032498"),
}
SELECTION_MAP = {"None": "Core Animals", "Core Males": "Core Animals",
                 "Core Females": "Core Animals"}
# STQ ORGAN_NAME → reference organ-weight column label.
ORGAN_MAP = {
    "Kidney-Left": "Kidney-Left",
    "Kidney-Right": "Kidney-Right",
    "Liver Weight": "Liver",
    "Terminal Bodyweight": "Terminal Body Weight",
}
ORGAN_COLS = ["Kidney-Left", "Kidney-Right", "Liver", "Terminal Body Weight"]


def _read(path: str) -> tuple[list[str], list[list]]:
    """Read the first (data) sheet of an STQ xlsx; return (header, rows).

    read_only is avoided — the STQ files, like NTP xlsx, can carry a bad stored
    dimension that makes read-only mode see one column."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    wb.close()
    header = [str(h) if h is not None else "" for h in rows[0]]
    return header, rows[1:]


def _sel(value) -> str:
    return SELECTION_MAP.get(str(value).strip(), "Core Animals")


def _cell(v) -> str:
    return "" if v is None else str(v)


def _write_csv(out_path: Path, header: list[str], rows: list[list]) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        for r in rows:
            w.writerow([_cell(c) for c in r])


def _sex_suffix(rows: list[list], sex_idx: int) -> str:
    sexes = {str(r[sex_idx]).strip().lower() for r in rows if r[sex_idx]}
    return sexes.pop() if len(sexes) == 1 else "all"


def convert_body_weight(path: str, outdir: Path) -> list[tuple[Path, str, int]]:
    header, rows = _read(path)
    idx = {c: header.index(c) for c in (
        "DEPOSITOR_STUDY_NUMBER", "DOSE", "ANMLNUM", "SEX", "SELECTION",
        "WGTVALUE", "OBSDAY", "TERMINAL_FLAG")}
    by_study: dict[str, list] = {}
    for r in rows:
        by_study.setdefault(r[idx["DEPOSITOR_STUDY_NUMBER"]], []).append(r)

    out_header = ["NTP Study Number", "Concentration", "Animal ID", "Sex",
                  "Selection", "Body Weight", "Observation Day", "Terminal Flag"]
    written = []
    for study, srows in by_study.items():
        chem, _ = STUDY_TO_CHEM.get(study, (study, None))
        study_out = _normalize_study(study)
        drows = [[
            study_out, r[idx["DOSE"]], _cell(r[idx["ANMLNUM"]]), r[idx["SEX"]],
            _sel(r[idx["SELECTION"]]), r[idx["WGTVALUE"]], r[idx["OBSDAY"]],
            "Yes" if str(r[idx["TERMINAL_FLAG"]]).strip() == "Yes" else "",
        ] for r in srows]
        sex = _sex_suffix(srows, idx["SEX"])
        out = outdir / chem / f"body_weight_truth_{sex}.csv"
        _write_csv(out, out_header, drows)
        written.append((out, chem, len(drows)))
    return written


def convert_organ_weight(path: str, outdir: Path) -> list[tuple[Path, str, int]]:
    header, rows = _read(path)
    idx = {c: header.index(c) for c in (
        "DEPOSITOR_STUDY_NUMBER", "DOSE", "ANMLID", "SEX", "SELECTION",
        "RMVLDAY", "ORGAN_NAME", "ORGANVALUE")}
    # One row per animal: gather each animal's organ rows into one record.
    animals: dict[tuple, dict] = {}
    order: list[tuple] = []
    for r in rows:
        study, anml = r[idx["DEPOSITOR_STUDY_NUMBER"]], _cell(r[idx["ANMLID"]])
        key = (study, anml)
        if key not in animals:
            animals[key] = {
                "study": study, "dose": r[idx["DOSE"]], "anml": anml,
                "sex": r[idx["SEX"]], "sel": r[idx["SELECTION"]],
                "rmvl": r[idx["RMVLDAY"]], "organs": {},
            }
            order.append(key)
        col = ORGAN_MAP.get(str(r[idx["ORGAN_NAME"]]).strip())
        if col:
            animals[key]["organs"][col] = r[idx["ORGANVALUE"]]

    by_study: dict[str, list] = {}
    for key in order:
        by_study.setdefault(animals[key]["study"], []).append(animals[key])

    out_header = ["NTP Study Number", "Concentration", "Animal ID", "Sex",
                  "Selection", "Removal Day"] + ORGAN_COLS
    written = []
    for study, alist in by_study.items():
        chem, _ = STUDY_TO_CHEM.get(study, (study, None))
        study_out = _normalize_study(study)
        drows = [
            [study_out, a["dose"], a["anml"], a["sex"], _sel(a["sel"]), a["rmvl"]]
            + [a["organs"].get(c, "") for c in ORGAN_COLS]
            for a in alist
        ]
        sexes = {str(a["sex"]).strip().lower() for a in alist if a["sex"]}
        sex = sexes.pop() if len(sexes) == 1 else "all"
        out = outdir / chem / f"organ_weights_truth_{sex}.csv"
        _write_csv(out, out_header, drows)
        written.append((out, chem, len(drows)))
    return written


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--body", default="assets/IAD_STQ_BODYWEIGHTS_BPAF_Triclosan_C08002-01_C03039-01.xlsx")
    ap.add_argument("--organ", default="assets/IAD_STQ_ORGANWEIGHTS_BPAF_Triclosan_C08002-01_C03039C-01.xlsx")
    ap.add_argument("--out", default="assets/converted")
    args = ap.parse_args()

    outdir = Path(args.out)
    print(f"writing NTP truth CSVs to {outdir}/<chemical>/")
    for out, chem, n in convert_body_weight(args.body, outdir):
        print(f"  body   {chem:10} {out.relative_to(outdir)}  ({n} rows)")
    for out, chem, n in convert_organ_weight(args.organ, outdir):
        print(f"  organ  {chem:10} {out.relative_to(outdir)}  ({n} animals)")


if __name__ == "__main__":
    main()
